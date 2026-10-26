"""Convert a blog-draft markdown file into a .docx for client review.

Reusable for any client draft written by the agency-blog-writer skill.
Handles headings (#..####), paragraphs, bullet and numbered lists,
blockquotes, **bold**, *italic*, [links](url) as real hyperlinks,
horizontal rules and simple pipe tables.

Usage: python scripts/md_to_docx.py [--landscape] draft.md [more.md ...]
Writes draft.docx next to each input. Use --landscape for wide
side-by-side tables (page-update docs). `inline code` renders in Consolas.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")


def add_hyperlink(par, text, url):
    part = par.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "005394")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_inline(par, text):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("[") and "](" in piece:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece).groups()
            if re.match(r"(https?:|mailto:|tel:)", url):
                add_hyperlink(par, label, url)
            else:
                # Repo-relative links mean nothing outside the repo; keep the text only
                par.add_run(label)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            par.add_run(piece[1:-1]).font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def convert(md_path, landscape=False):
    md_path = Path(md_path)
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    if landscape:
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(0.6))
        doc.styles["Normal"].font.size = Pt(10)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", s):
            doc.add_paragraph("_" * 40)
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip().strip("|").replace(r"\|", "\uE000")
                cells = [c.strip().replace("\uE000", "|") for c in row.split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, val in enumerate(cells):
                    cell = table.cell(r, c)
                    # "<br>" in a cell starts a new paragraph; a leading "- " becomes a bullet
                    for k, part in enumerate(re.split(r"<br\s*/?>", val)):
                        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                        if bm := re.match(r"[-*]\s+(.*)", part.strip()):
                            p.style = "List Bullet"
                            part = bm.group(1)
                        add_inline(p, part.strip())
                        if r == 0:
                            for run in p.runs:
                                run.bold = True
            # Repeat the header row on each page, and size columns by how much text they hold
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
            ncols = len(table.columns)
            weights = [
                max(1.0, sum(len(row[c]) for row in rows if c < len(row)) / len(rows)) ** 0.5
                for c in range(ncols)
            ]
            usable = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
            table.autofit = False
            for c in range(ncols):
                width = int(usable * weights[c] / sum(weights))
                table.columns[c].width = width
                for row in table.rows:
                    row.cells[c].width = width
            continue
        elif m := re.match(r"(#{1,4})\s+(.*)", s):
            level = len(m.group(1))
            h = doc.add_heading(level=level)
            add_inline(h, m.group(2))
        elif s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, s.lstrip("> ").strip())
        elif m := re.match(r"[-*+]\s+(.*)", s):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
        elif m := re.match(r"\d+[.)]\s+(.*)", s):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(1))
        elif s.startswith("[IMAGE"):
            p = doc.add_paragraph()
            run = p.add_run(s)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            p = doc.add_paragraph()
            add_inline(p, s)
        i += 1
    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    args = sys.argv[1:]
    landscape = "--landscape" in args
    for arg in (a for a in args if a != "--landscape"):
        print(convert(arg, landscape=landscape))
