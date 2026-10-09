"""Build book-schema-review.xlsx from titles.json (same data as build_schema.py).

Reusable: run after refresh_titles.py / build_schema.py:  python -I build_review_sheet.py
Tabs: Read Me, Titles (one row per textbook), Variants (one row per edition/format), JSON-LD (the code).
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import sys
sys.path.insert(0, str(Path(__file__).parent))
import build_schema as bs  # reuses graph(); also regenerates the .html files (idempotent)

HERE = Path(__file__).parent
titles = json.load(open(HERE / "titles.json", encoding="utf-8"))

F = "Arial"
HDR = Font(name=F, bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1F4E78")
BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
INPUT = PatternFill("solid", fgColor="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top"); TOP = Alignment(vertical="top")
thin = Side(style="thin", color="D9D9D9"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def table(ws, headers, rows, widths, review_cols=()):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, HFILL, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = BODY, WRAP, BOX
            if headers[c.column - 1] in review_cols:
                c.fill = INPUT
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32


def yn(ws, col_letter, n):
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{n + 1}")


wb = Workbook()

# Read Me
rm = wb.active; rm.title = "Read Me"
lines = [
    ("Mujo Book schema: review sheet", BOLD),
    ("Built from clients/mujo/schema/titles.json (live WooCommerce data pulled 2026-10-10, authors confirmed by Brittni 2026-10-09). Nothing here is live on mujo.com yet.", BODY),
    ("", BODY),
    ("How to review", BOLD),
    ("1. Titles tab: one row per textbook. Check book title, authors, edition and description. Mark 'OK?' Y or N and add a comment if anything is wrong.", BODY),
    ("2. Variants tab: one row per edition/format (128 rows). Check ISBN, format, price and name. Mark 'OK?' Y or N.", BODY),
    ("3. JSON-LD tab: the exact code per title, for whoever adds it to the site. No need to review line by line.", BODY),
    ("Only edit the yellow columns (OK? and Comments). Fixes go back into titles.json / refresh_titles.py, then the files are regenerated.", BODY),
    ("", BODY),
    ("Example of a filled-in review row", BOLD),
    ("OK? = N  |  Comments = 'Edition should be 2nd edition, not 1st.'", BODY),
    ("", BODY),
    ("Sources", BOLD),
    ("Prices, ISBNs (GTINs), SKUs and variant URLs: live mujo.com product pages (Rank Math ProductGroup output). Authors: VitalSource product pages, confirmed or corrected by Brittni. Editions: VitalSource (higher ed only).", BODY),
    ("Publisher: 'Mujo Learning Systems' (Brittni, 2026-10-09). Validated pattern: Google Rich Results Test 2026-10-10, 0 errors.", BODY),
]
for text, font in lines:
    rm.append([text]); rm.cell(rm.max_row, 1).font = font; rm.cell(rm.max_row, 1).alignment = WRAP
rm.column_dimensions["A"].width = 120

# Titles
ws = wb.create_sheet("Titles")
rows = []
for t in titles:
    rows.append([t["product_id"], "Higher ed" if t["audience"] == "higher-ed" else "High school CTE", t["book"],
                 ", ".join(t["authors"]), t["edition"] or "(not set)", "Mujo Learning Systems", len(t["variants"]),
                 t["name"], t["desc"], t["url"], t["notes"], None, None])
hdr = ["Product ID", "Audience", "Book title (schema)", "Author(s)", "Edition", "Publisher", "# Variants",
       "Product group name", "Description", "Page URL", "Notes from build", "OK?", "Comments"]
table(ws, hdr, rows, [10, 14, 34, 26, 11, 20, 9, 40, 60, 45, 40, 7, 40], review_cols=("OK?", "Comments"))
yn(ws, "L", len(rows))

# Variants
wv = wb.create_sheet("Variants")
rows = []
for t in titles:
    for v in t["variants"]:
        rows.append([t["product_id"], t["book"], f"{t['book']} - {v['label']}", v["label"],
                     "eBook" if v["format"] == "EBook" else "Print (Paperback)", v["isbn"], v["sku"],
                     float(v["price"]), ", ".join(t["authors"]), v["url"], v["img"], None, None])
hdr = ["Product ID", "Book title", "Variant name (schema)", "Edition / license", "Format", "ISBN", "SKU",
       "Price (USD)", "Author(s)", "Offer URL", "Image", "OK?", "Comments"]
table(wv, hdr, rows, [10, 30, 46, 34, 14, 15, 22, 10, 24, 50, 50, 7, 40], review_cols=("OK?", "Comments"))
for c in wv["H"][1:]:
    c.number_format = '$#,##0.00'
for c in wv["F"][1:]:
    c.number_format = "@"
yn(wv, "L", len(rows))

# JSON-LD
wj = wb.create_sheet("JSON-LD")
rows = [[t["product_id"], t["book"], t["url"], json.dumps(bs.graph(t), indent=1, ensure_ascii=False)] for t in titles]
table(wj, ["Product ID", "Book title", "Page URL", "JSON-LD (paste-ready)"], rows, [10, 30, 45, 110])
for c in wj["D"][1:]:
    c.font = Font(name="Consolas", size=9)

out = HERE / "book-schema-review.xlsx"
wb.save(out)
print("saved", out, len(titles), "titles", sum(len(t["variants"]) for t in titles), "variants")
