import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
d = json.load(open(src))

# VitalSource lookups, 2026-10-08 (keyed by student eBook ISBN)
vs = {
 "9781998671892": ("AI Literacy for Healthcare, 1st edition", "David Shaw", "https://www.vitalsource.com/products/ai-literacy-for-healthcare-david-shaw-v9781998671892"),
 "9781998671663": ("Artificial Intelligence Business Analytics, 1st edition", "Katrina Garofalo", "https://www.vitalsource.com/products/artificial-intelligence-business-analytics-katrina-garofalo-v9781998671663"),
 "9781998671014": ("Artificial Intelligence Project Management, 1st edition", "David Shaw", "https://www.vitalsource.com/products/artificial-intelligence-project-management-david-shaw-v9781998671014"),
 "9781998798704": ("Artificial Intelligence Literacy, 1st edition", "David Shaw", "https://www.vitalsource.com/products/artificial-intelligence-literacy-david-shaw-v9781998798704"),
 "9781998798780": ("Artificial Intelligence for Entrepreneurs, 1st edition", "David Shaw", "https://www.vitalsource.com/products/artificial-intelligence-for-entrepreneurs-david-shaw-v9781998798780"),
 "9781998798940": ("Artificial Intelligence Business Administration, 1st edition", "David Shaw", "https://www.vitalsource.com/products/artificial-intelligence-business-administration-david-shaw-v9781998798940"),
 "9781998798865": ("Artificial Intelligence and Human Resource Management, 1st edition", "Katrina Garofalo", "https://www.vitalsource.com/products/artificial-intelligence-and-human-resource-katrina-garofalo-v9781998798865"),
 "9781998798230": ("Artificial Intelligence Marketing, 1st edition", "David Shaw", "https://www.vitalsource.com/products/artificial-intelligence-marketing-david-shaw-v9781998798230"),
 "9781998798742": ("Generative AI for Business, 1st edition", "David Shaw", "https://www.vitalsource.com/products/generative-ai-for-business-david-shaw-v9781998798742"),
 "9781998671137": ("Artificial Intelligence Business Writing, 1st edition", "Katrina Garofalo", "https://www.vitalsource.com/products/artificial-intelligence-business-writing-katrina-garofalo-v9781998671137"),
 "9781998671175": ("User Interface/User Experience Design, 1st edition", "Katrina Garofalo", "https://www.vitalsource.com/products/user-interface-user-experience-design-katrina-garofalo-v9781998671175"),
 "9781988940809": ("Pay-Per-Click Advertising, 1st edition", "David Shaw", "https://www.vitalsource.com/products/pay-per-click-advertising-david-shaw-v9781988940809"),
 "9781988940656": ("Strategic Web Design & e-Commerce, 1st edition", "Shawn Moore", "https://www.vitalsource.com/products/strategic-web-design-and-e-commerce-shawn-moore-amp-adam-wilkins-v9781988940656"),
 "9781998798315": ("Public Relations Strategy and Communications, 1st edition", "Emma Hatfield", "https://www.vitalsource.com/products/public-relations-and-crisis-management-emma-hatfield-v9781998798315"),
 "9781998798001": ("Big Data Analytics and Reporting, 1st edition", "David Shaw", "https://www.vitalsource.com/products/big-data-analytics-reporting-david-shaw-v9781998798001"),
 "9781988940885": ("Writing Digital Media Content, 1st edition", "Rebecca Saloustros", "https://www.vitalsource.com/products/writing-digital-media-content-student-39-s-manual-rebecca-saloustros-v9781988940885"),
 "9781988940694": ("Customer Relationship Management Marketing Automation, 1st edition", "Emma Hatfield", "https://www.vitalsource.com/products/customer-relationship-management-marketing-emma-hatfield-v9781988940694"),
 "9781998798278": ("Influencer Marketing Fundamentals, 1st edition", "David Shaw", "https://www.vitalsource.com/products/influencer-marketing-fundamentals-david-shaw-v9781998798278"),
 "9781998798155": ("Search Engine Optimization, 4th edition", "David Shaw", "https://www.vitalsource.com/products/search-engine-optimization-david-shaw-v9781998798155"),
 "9781998671335": ("Social Media Marketing Strategies, 7th edition", "David Shaw", "https://www.vitalsource.com/products/social-media-marketing-strategies-david-shaw-v9781998671335"),
 "9781998798032": ("Principles of Marketing, 1st edition", "David Shaw, Alex Strauss", "https://www.vitalsource.com/products/principles-of-marketing-david-shaw-v9781998798032"),
 "9781998798490": ("Digital Marketing Fundamentals, 3rd edition", "Shaw & Wilkins", "https://www.vitalsource.com/products/digital-marketing-fundamentals-shaw-amp-wilkins-v9781998798490"),
 "9781998671700": ("Prompt Engineering and LLMs, 1st edition", "David Shaw", "https://www.vitalsource.com/products/prompt-engineering-and-llms-david-shaw-v9781998671700"),
 "9781988940519": ("Website Design Strategy, 1st edition", "Shawn Moore", "https://www.vitalsource.com/products/website-design-strategy-adam-wilkins-v9781988940519"),
 "9781988940472": ("Online Marketing Fundamentals: High School Edition, 1st edition", "Shawn Moore", "https://www.vitalsource.com/products/online-marketing-fundamentals-adam-wilkins-v9781988940472"),
 "9781988940496": ("Creating Digital Media Content: High School Edition, 1st edition", "Rebecca Saloustros", "https://www.vitalsource.com/products/writing-online-content-rebecca-saloustros-v9781988940496"),
}
# Student eBook ISBNs with no VitalSource listing; every other ISBN for the title was also checked (2026-10-08)
not_on_vs = {"9781998671588","9781998671625","9781998798988","9781998671304","9781998798650","9781998798520",
             "9781998798445","9781998671236","9781998798414","9781998671441","9781998671434","9781998671465"}
notes = {
 "9781988940656": "VitalSource text lists Shawn Moore only, but its URL also names Adam Wilkins. Confirm.",
 "9781988940519": "VitalSource text lists Shawn Moore, but its URL names Adam Wilkins. Confirm. VitalSource title differs: 'Website Design Strategy'.",
 "9781988940472": "VitalSource text lists Shawn Moore, but its URL names Adam Wilkins. Confirm.",
 "9781998798490": "VitalSource gives surnames only ('Shaw & Wilkins'). Probably David Shaw & Adam Wilkins. Confirm full names.",
 "9781998798032": "Two authors: David Shaw and Alex Strauss (Mujo's President).",
 "9781998798155": "4th edition (not 1st).",
 "9781998671335": "7th edition (not 1st).",
 "9781998798315": "VitalSource URL says 'public-relations-and-crisis-management' (older title?).",
 "9781988940496": "VitalSource URL says 'writing-online-content' (older title?).",
}
title_flags = {
 27226: "Instructor ISBN 9781988798285 fails the ISBN check digit. Likely typo for 9781998798285 (978-1-998... prefix like the rest). Fix the GTIN in WooCommerce.",
 27146: "Variant SKUs look swapped: 'AIMF-Resource-Print' is the eBook resource ($398), plain 'AIMF' is the print resource ($99).",
 27125: "Resource print is $497 here vs $99 on most CTE titles. Confirm.",
 27118: "Resource print is $497 here vs $99 on most CTE titles. Confirm.",
}

def ok(s): return len(s)==13 and s.isdigit() and sum(int(c)*(1 if i%2==0 else 3) for i,c in enumerate(s))%10==0

hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F4E78")
warn_fill = PatternFill("solid", fgColor="FFF2CC"); good_fill = PatternFill("solid", fgColor="E2EFDA")

wb = Workbook()
ws = wb.active; ws.title = "Titles"
cols = ["Audience","Product ID","Title on mujo.com","Page URL","Student eBook ISBN","Student print ISBN","Instructor/Resource eBook ISBN","Resource print ISBN",
        "Author(s) (VitalSource)","VitalSource title","VitalSource link","Status","Britt confirms author (Y/N)","Notes"]
ws.append(cols)
for a,pid,name,url,v in d:
    stud_e = next((g for s,g,p,e,f in v if f=="ebook" and e in ("for-students","for-teachers-student")), "")
    stud_p = next((g for s,g,p,e,f in v if f=="print" and e=="for-teachers-student"), "")
    res_e = next((g for s,g,p,e,f in v if f=="ebook" and e in ("for-instructors","for-teachers-resource")), "")
    res_p = next((g for s,g,p,e,f in v if f=="print" and e=="for-teachers-resource"), "")
    t, au, link = vs.get(stud_e, ("","",""))
    status = "Found on VitalSource" if au else ("Not on VitalSource (all ISBNs checked). Ask Britt" if stud_e in not_on_vs else "Pending lookup")
    n = "; ".join(x for x in (notes.get(stud_e,""), title_flags.get(pid,"")) if x)
    ws.append([a,pid,name,url,stud_e,stud_p,res_e,res_p,au,t,link,status,"",n])
    r = ws.max_row
    ws.cell(r,12).fill = good_fill if au else warn_fill
    if pid in title_flags: ws.cell(r,14).fill = warn_fill

ws2 = wb.create_sheet("All variants")
ws2.append(["Product ID","Title","SKU","ISBN / GTIN","Valid ISBN-13","Price (USD)","Edition","Format"])
for a,pid,name,url,v in d:
    for s,g,p,e,f in v:
        ws2.append([pid,name,s,g,"Yes" if ok(g) else "NO, check digit fails",float(p),e,f])
        if not ok(g):
            for c in range(1,9): ws2.cell(ws2.max_row,c).fill = warn_fill

ws3 = wb.create_sheet("About")
for line in [
 "Mujo textbook authors and ISBNs, built 2026-10-08 by ThinkProfits (Francis + Claude) for Book schema work.",
 "Sources: ISBNs = GTINs on mujo.com product variants (Rank Math ProductGroup JSON-LD). Authors = VitalSource product pages, looked up by student eBook ISBN.",
 "38 textbook titles on the live shop (25 higher ed, 13 high school CTE). Bundles/programs and unpublished Texas Edition drafts are excluded.",
 "Instructor/teacher-resource ISBNs mostly aren't listed on VitalSource, so authors are looked up by the student eBook ISBN.",
 "12 titles aren't on VitalSource under any of their ISBNs (checked 2026-10-08). Britt to fill in those authors. VitalSource rate-limits fast lookups: go one every ~6 seconds.",
 "Edition: most are 1st edition, but SEO is 4th, Social Media Marketing Strategies 7th, Digital Marketing Fundamentals 3rd. Use the real edition in schema.",
 "Use in schema: author = Person (named author), publisher = Mujo Learning Systems Inc. Never use Mujo as author. See memory/mujo-context.md.",
 "Britt to confirm column M before any author goes live.",
]: ws3.append([line])
ws3.column_dimensions["A"].width = 140

for sh in (ws, ws2):
    for c in sh[1]:
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="top")
    sh.freeze_panes = "A2"; sh.auto_filter.ref = sh.dimensions
    for i,col in enumerate(sh.columns,1):
        w = max(len(str(c.value or "")) for c in col)
        sh.column_dimensions[get_column_letter(i)].width = min(max(12, w+2), 60)
for row in ws.iter_rows(min_row=2):
    for c in row:
        if c.column_letter in ("E","F","G","H"): c.number_format = "@"
for row in ws2.iter_rows(min_row=2):
    row[3].number_format = "@"; row[5].number_format = '"$"#,##0'
wb.save(out)
print("saved", out, ws.max_row-1, "titles,", sum(1 for r in ws.iter_rows(min_row=2) if r[8].value), "with authors")
